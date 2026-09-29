#!/usr/bin/env python3
"""results/summary.csv -> summary.xlsx, highlighting where MATCHA beats every baseline.

    python results/utils/build_summary_xlsx.py

A MATCHA cell (Human-led or LLM-led) is highlighted when it is STRICTLY greater than
all four baselines available to that model on that dataset:

    Human-only      dataset-level, no routing, no uncertainty
    LLM-only        model-level
    Crowd-LLM       model-level
    CoAnnotating    this row's uncertainty basis

Human-only / LLM-only / Crowd-LLM are written once on the first row of each dataset
block because they do not vary with the uncertainty basis; they are carried down and
applied to all three rows of the block, so a MATCHA value on the Entropy row is still
compared against them.

STRICTLY greater: a tie is not an outperformance and is left unhighlighted. Comparison
is at 4 decimal places -- the precision the table displays -- because the underlying
floats differ in the last digits between sources (crowd_aggregation.csv stores full
precision, routing.csv rounds to 4dp), which would otherwise turn exact display-level
ties into spurious wins.

Human-led = h-human (routed instances take the crowd label, the LLM label is discarded)
LLM-led   = h-human_llm (the LLM stays in the annotator pool on routed instances)

Writes results/summary.xlsx
"""
import csv
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[2]
RES = ROOT / "results"

DATASETS = {"Sentiment Polarity", "MovieReviews", "CrowdTruth RelEx",
            "CoNLL-2003 NER", "PICO", "QUIZ", "LabelMe", "ImageNet-16H"}
BASES = {"Self-Eval.", "Entropy", "Inter-rater"}
# per model block: (LLM-only, Crowd-LLM, basis-label, CoAnnotating, MATCHA-HL, MATCHA-LL)
BLOCKS = [(3, 4, 6, 7, 8, 9), (10, 11, 13, 14, 15, 16), (17, 18, 20, 21, 22, 23)]
HUMAN_ONLY = 2

WIN = PatternFill("solid", fgColor="C6EFCE")     # green fill, Excel's standard "good"
WIN_FONT = Font(bold=True, color="006100")
HDR = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")


def main():
    rows = [list(r) + [""] * (24 - len(r))
            for r in csv.reader(open(RES / "summary.csv", encoding="utf-8-sig"))]

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"

    carried = {}          # per model block: (human_only, llm_only, crowd_llm)
    n_win = 0
    for i, r in enumerate(rows, start=1):
        name = r[1].strip()
        is_hdr = name in ("Dataset",) or r[0].strip() in ("Modality",) or (
            r[3].strip() in ("LLM-only",) or r[8].strip() in ("Human-led",))
        if name in DATASETS:                       # first row of a block: refresh
            for bi, (lo, cl, _, _, _, _) in enumerate(BLOCKS):
                carried[bi] = (r[HUMAN_ONLY].strip(), r[lo].strip(), r[cl].strip())

        for c, val in enumerate(r, start=1):
            cell = ws.cell(row=i, column=c, value=_num(val))
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if is_hdr:
                cell.fill = HDR
                cell.font = Font(bold=True)

        if is_hdr or r[6].strip() not in BASES:
            continue
        for bi, (lo, cl, _, ca, hl, ll) in enumerate(BLOCKS):
            base = [_f(x) for x in (*carried.get(bi, ("", "", "")), r[ca])]
            base = [b for b in base if b is not None]
            if not base:
                continue
            top = max(base)
            for col in (hl, ll):
                v = _f(r[col])
                if v is not None and round(v, 4) > round(top, 4):
                    ws.cell(row=i, column=col + 1).fill = WIN
                    ws.cell(row=i, column=col + 1).font = WIN_FONT
                    n_win += 1

    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 19
    for c in range(3, 25):
        ws.column_dimensions[get_column_letter(c)].width = 11.5
    ws.freeze_panes = "C4"

    out = RES / "summary.xlsx"
    wb.save(out)
    print(f"  {out.relative_to(ROOT)}")
    print(f"  MATCHA cells highlighted (strictly beat all baselines): {n_win}")


def _f(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def _num(s):
    v = _f(s)
    return v if v is not None else s


if __name__ == "__main__":
    main()
